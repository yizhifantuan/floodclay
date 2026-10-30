#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""从两个原始 Excel 文件生成 19 地区卫星拍摄时间轴。

安装：python -m pip install -r requirements.txt
运行：python plot_timelines.py
也可使用 --dates / --events / --output / --dpi / --font 指定路径和样式。
"""

from __future__ import annotations

import argparse
import csv
import math
import re
from datetime import datetime, timedelta
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.dates as mdates
import matplotlib.pyplot as plt
from matplotlib import font_manager
from matplotlib.ticker import FixedLocator
from openpyxl import load_workbook
from PIL import Image, ImageOps


BASE = Path(__file__).resolve().parent
DATE_FILE = "19地区_S1_S2_PlanetScope_拍摄日期汇总.xlsx"
EVENT_FILE = "洪涝事件.xlsx"
PLATFORMS = ("S1", "S2", "PlanetScope")
COLORS = {"S1": "#3862EF", "S2": "#1DA64D", "PlanetScope": "#8039ED"}
BACKGROUND = "#FCFCFD"


def configure_font(font_path: str | None) -> None:
    """优先使用中文字体；SVG 将文字转为路径，便于跨电脑使用。"""
    if font_path:
        path = Path(font_path)
        if not path.is_file():
            raise FileNotFoundError(path)
        font_manager.fontManager.addfont(str(path))
        family = font_manager.FontProperties(fname=str(path)).get_name()
    else:
        names = {item.name for item in font_manager.fontManager.ttflist}
        preferred = ["Microsoft YaHei", "SimHei", "Noto Sans CJK SC", "Source Han Sans SC", "PingFang SC", "WenQuanYi Zen Hei", "Arial Unicode MS"]
        family = next((name for name in preferred if name in names), None)
        if family is None:
            raise RuntimeError("未找到中文字体。请安装微软雅黑/思源黑体，或用 --font 指定字体文件。")
    plt.rcParams.update({
        "font.family": family,
        "axes.unicode_minus": False,
        "svg.fonttype": "path",
        "font.size": 12,
        "figure.facecolor": BACKGROUND,
        "axes.facecolor": BACKGROUND,
    })


def parse_period(text: str) -> tuple[datetime, datetime]:
    """解析 2017/03/14–05/02，支持完整结束年份和跨年简写。"""
    match = re.fullmatch(
        r"\s*(\d{4})[/-](\d{1,2})[/-](\d{1,2})\s*[–—~～至-]\s*(?:(\d{4})[/-])?(\d{1,2})[/-](\d{1,2})\s*",
        str(text),
    )
    if not match:
        raise ValueError(f"无法解析洪水期：{text!r}")
    year, month, day, end_year, end_month, end_day = match.groups()
    start = datetime(int(year), int(month), int(day))
    end = datetime(int(end_year or year), int(end_month), int(end_day))
    if end < start and end_year is None:
        end = end.replace(year=start.year + 1)
    if end < start:
        raise ValueError(f"洪水期结束早于开始：{text}")
    return start, end


def read_dates(path: Path) -> dict:
    """使用日期列表页，不使用地区汇总页中损坏的统计公式。"""
    workbook = load_workbook(path, read_only=True, data_only=True)
    regions: dict = {}
    try:
        for row in workbook["日期列表"].iter_rows(min_row=5, values_only=True):
            chinese, english, platform, raw = row[:4]
            if not any(value is not None for value in row[:4]):
                continue
            if not all((chinese, english, platform, raw)):
                raise ValueError(f"拍摄日期记录不完整：{row[:4]}")
            key = str(english).strip()
            platform = str(platform).strip()
            if platform not in PLATFORMS:
                raise ValueError(f"未知卫星平台：{platform}")
            region = regions.setdefault(key, {"chinese": str(chinese).strip(), "dates": {}})
            if platform in region["dates"]:
                raise ValueError(f"重复地区平台记录：{key}/{platform}")
            # 原表同时包含中文分号、换行，以及字面量反斜杠 n。
            tokens = re.findall(r"\d{4}-\d{2}-\d{2}", str(raw))
            residue = re.sub(r"\d{4}-\d{2}-\d{2}", "", str(raw)).replace("\\n", "")
            if not tokens or re.sub(r"[\s；;,，、]", "", residue):
                raise ValueError(f"日期列表含无法解析的内容：{key}/{platform}: {raw}")
            parsed = [datetime.strptime(token, "%Y-%m-%d") for token in tokens]
            if len(parsed) != len(set(parsed)):
                raise ValueError(f"同一地区平台有重复日期：{key}/{platform}")
            region["dates"][platform] = sorted(parsed)
    finally:
        workbook.close()
    if len(regions) != 19:
        raise ValueError(f"预期 19 地区，实际读取 {len(regions)} 个")
    for key, region in regions.items():
        if set(region["dates"]) != set(PLATFORMS):
            raise ValueError(f"地区缺少平台日期：{key}")
    return regions


def read_events(path: Path) -> dict:
    workbook = load_workbook(path, read_only=True, data_only=True)
    events = {}
    try:
        for row in workbook["洪涝事件"].iter_rows(min_row=2, values_only=True):
            label, period = row[:2]
            # 表尾有仅写地区名但没有事件日期的零散备注，不作为事件。
            if not period:
                continue
            match = re.search(r"([A-Za-z][A-Za-z ]*)\s*$", str(label))
            if not match:
                raise ValueError(f"缺少英文地区键：{label}")
            key = match.group(1).strip()
            if key in events:
                raise ValueError(f"重复事件：{key}")
            start, end = parse_period(str(period))
            events[key] = {"start": start, "end": end, "cause": row[2]}
    finally:
        workbook.close()
    return events


def draw_region(index: int, key: str, region: dict, event: dict, output: Path, dpi: int) -> tuple[Path, list]:
    dates = region["dates"]
    start, end = event["start"], event["end"]
    all_dates = [date for values in dates.values() for date in values]
    earliest, latest = min([start, *all_dates]), max([end, *all_dates])
    span = max(1, (latest - earliest).days)
    margin = max(1.2, span * 0.045)
    left, right = earliest - timedelta(days=margin), latest + timedelta(days=margin)

    figure, ax = plt.subplots(figsize=(12, 4.6))
    figure.subplots_adjust(left=0.105, right=0.968, bottom=0.24, top=0.73)
    figure.text(0.052, 0.925, f"{index:02d}  {region['chinese']}  {key}",
                ha="left", va="center", fontsize=19, color="#20252D")
    # 日期只有日精度：每个日期以当天 00:00 为中心，阴影两端各扩半天，包含首尾日期。
    ax.axvspan(start - timedelta(hours=12), end + timedelta(hours=12),
               ymin=0.025, ymax=0.975, facecolor="#F4BF4F", alpha=0.18, linewidth=0, zorder=0)
    center = start + (end - start) / 2
    ax.text(center, 1.20, "洪水期", transform=ax.get_xaxis_transform(),
            ha="center", va="center", fontsize=15, color="#D67A00", clip_on=False)
    ax.text(center, 1.075, f"{start:%Y-%m-%d} 至 {end:%Y-%m-%d}",
            transform=ax.get_xaxis_transform(), ha="center", va="center",
            fontsize=10, color="#A36C20", clip_on=False)

    records = []
    for platform, height in zip(PLATFORMS, [2, 1, 0]):
        ax.hlines(height, left, right, color="#CBCDCE", linewidth=2.7, zorder=1)
        ax.scatter(dates[platform], [height] * len(dates[platform]), s=104,
                   color=COLORS[platform], edgecolors="none", zorder=3)
        records.extend([index, region["chinese"], key, platform, date.strftime("%Y-%m-%d"),
                        start.strftime("%Y-%m-%d"), end.strftime("%Y-%m-%d"),
                        "是" if start <= date <= end else "否"] for date in dates[platform])

    step = max(1, math.ceil(span / 8))
    ticks = [earliest + timedelta(days=offset) for offset in range(0, span + 1, step)]
    ax.xaxis.set_major_locator(FixedLocator(mdates.date2num(ticks)))
    ax.xaxis.set_major_formatter(mdates.DateFormatter("%m-%d"))
    ax.set_xlim(left, right)
    ax.set_ylim(-0.5, 2.5)
    ax.set_yticks([2, 1, 0], labels=["S1", "S2", "PS"])
    ax.tick_params(axis="y", length=0, pad=25, labelsize=19, colors="#22262C")
    ax.tick_params(axis="x", length=0, pad=12, labelsize=12, colors="#353A42")
    ax.set_axisbelow(True)
    ax.grid(axis="x", color="#DCDDE0", linestyle=(0, (3, 3)), linewidth=1.0)
    for spine in ax.spines.values():
        spine.set_visible(False)
    years = str(earliest.year) if earliest.year == latest.year else f"{earliest.year}–{latest.year}"
    ax.set_xlabel(f"拍摄日期（{years} 年）", labelpad=13, fontsize=11, color="#636A73")
    figure.text(0.052, 0.047, "S1：Sentinel-1    S2：Sentinel-2    PS：PlanetScope",
                fontsize=9, color="#777E87")
    figure.text(0.968, 0.047,
                f"拍摄日期数：S1 {len(dates['S1'])}    S2 {len(dates['S2'])}    PS {len(dates['PlanetScope'])}",
                ha="right", fontsize=9, color="#777E87")

    filename = f"{index:02d}_{key}_{region['chinese']}"
    png = output / "PNG" / f"{filename}.png"
    figure.savefig(png, dpi=dpi, facecolor=BACKGROUND)
    figure.savefig(output / "SVG" / f"{filename}.svg", facecolor=BACKGROUND)
    plt.close(figure)
    return png, records


def create_overview(images: list[Path], output: Path) -> None:
    """只作浏览用缩略预览，正式图为 PNG 文件夹内的 19 张单图。"""
    width, height = 900, 345
    gap, columns = 14, 3
    rows = math.ceil(len(images) / columns)
    overview = Image.new("RGB", (columns * width + (columns + 1) * gap,
                                 rows * height + (rows + 1) * gap), "#E6E8EC")
    for index, path in enumerate(images):
        with Image.open(path) as image:
            tile = ImageOps.contain(image.convert("RGB"), (width, height), Image.Resampling.LANCZOS)
            overview.paste(tile, (gap + (index % columns) * (width + gap),
                                  gap + (index // columns) * (height + gap)))
    overview.save(output / "19地区_总览预览.jpg", quality=94)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--dates", type=Path, default=BASE / "inputs" / DATE_FILE)
    parser.add_argument("--events", type=Path, default=BASE / "inputs" / EVENT_FILE)
    parser.add_argument("--output", type=Path, default=BASE / "figures")
    parser.add_argument("--dpi", type=int, default=300)
    parser.add_argument("--font", help="可选的中文字体文件路径，如 C:/Windows/Fonts/msyh.ttc")
    args = parser.parse_args()
    if args.dpi < 72:
        parser.error("--dpi 请至少设置为 72")
    configure_font(args.font)
    regions, events = read_dates(args.dates), read_events(args.events)
    if set(regions) != set(events):
        raise ValueError(f"两个表的地区不一致：{set(regions) ^ set(events)}")
    for subdirectory in ("PNG", "SVG"):
        (args.output / subdirectory).mkdir(parents=True, exist_ok=True)
    images, records, summaries = [], [], []
    for index, (key, region) in enumerate(regions.items(), start=1):
        image, entries = draw_region(index, key, region, events[key], args.output, args.dpi)
        images.append(image)
        records.extend(entries)
        counts = [len(region["dates"][platform]) for platform in PLATFORMS]
        summaries.append([index, region["chinese"], key, events[key]["start"].strftime("%Y-%m-%d"),
                          events[key]["end"].strftime("%Y-%m-%d"), *counts, sum(counts)])
        print(f"{index:02d}/19 {key}: S1={counts[0]}, S2={counts[1]}, PS={counts[2]}")
    with (args.output / "绘图数据.csv").open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["序号", "地区中文", "地区英文", "平台", "拍摄日期", "洪水开始", "洪水结束", "是否洪水期内"])
        writer.writerows(records)
    with (args.output / "地区核对.csv").open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["序号", "地区中文", "地区英文", "洪水开始", "洪水结束", "S1日期数", "S2日期数", "PS日期数", "日期总数"])
        writer.writerows(summaries)
    create_overview(images, args.output)
    print(f"完成：19 张 PNG（{args.dpi} dpi），19 张 SVG；共 {len(records)} 个地区/平台/日期点。")


if __name__ == "__main__":
    main()
